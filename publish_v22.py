"""Repository report pipeline: preserve dated reports; append v2.2 evidence."""
import json, shutil, hashlib
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
from collections import defaultdict
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from official_data_v22 import collect


def save():
    site=Path('docs');q=json.loads((site/'quant_v22.json').read_text());a=json.loads((site/'reliability.json').read_text());d=json.loads((site/'data.json').read_text())
    if len({q['asof'],a['asof'],d['date']})!=1:raise ValueError('asof mismatch')
    official=collect(output=site/'official_v22.json')
    groups=defaultdict(list);signals={s['code']:s for s in d['stocks']}
    for s in official.get('dart',{}).get('stocks',[]):
        sector=s.get('industry_code_current');val=signals.get(s['code'],{}).get('price60')
        if sector and isinstance(val,(int,float)):groups[sector].append((s['code'],val))
    for s in official.get('dart',{}).get('stocks',[]):
        peers=groups.get(s.get('industry_code_current'),[]);v=signals.get(s['code'],{}).get('price60')
        others=[x for c,x in peers if c!=s['code']]
        s['industry_relative']={'status':'exploratory' if len(others)>=3 and v is not None else 'unconfirmed','peer_n':len(others),'return60_minus_peer_mean_pct':round(v-sum(others)/len(others),2) if len(others)>=3 and v is not None else None,'basis':'current DART industry code, candidate-universe equal-weight peers; NOT official KRX industry index or historical membership'}
    (site/'official_v22.json').write_text(json.dumps(official,ensure_ascii=False,indent=2))
    # Reuse the project's existing workbook pipeline and styles.
    reportdir=site/'reports';before={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in reportdir.glob('*.xlsx') if 'latest' not in p.name}
    wb=load_workbook(reportdir/'latest.xlsx')
    def sheet(name,headers):
        if name in wb:del wb[name]
        ws=wb.create_sheet(name);ws.append(headers);ws.freeze_panes='C2';return ws
    ws=sheet('v22_품질감사',['코드','종목','기존제외','원본봉수','정상연속봉수','격리봉수','가격검증','구간회복','독립대조','원인분류','남은오류','NAVER 원본 SHA256','YAHOO 원본 SHA256'])
    for r in a['stocks']:ws.append([r['code'],r['name'],r['v21_excluded'],r.get('rows'),r.get('usable_segment_rows'),r.get('quarantined_rows'),r['strict_valid'],r.get('recovered'),r.get('cross_validation',{}).get('status'),', '.join(r.get('cause_categories',[])),r.get('strict_error') or r.get('error'),r.get('raw_sha256'),r.get('yahoo_raw_sha256')])
    ws=sheet('v22_RSI比較',['코드','종목','기간','조건','상태','표본','종목수','분기수','손절선도달(%)','평균MAE(%)','평균순수익(%)','손익비','전략승률(%)','기간말상승빈도(%)','하락보합빈도(%)','목표선도달(%)','평균MFE(%)','평균유사거리','연도분포','500건평균(%)','무제한평균(%)'])
    for s in q['stocks']:
        for h,b in s.get('periods',{}).items():
            for label,p in [('기존조건',b),('RSI추가',s.get('rsi_periods',{}).get(h,{}))]:
                ws.append([s['code'],s['name'],h,label,'연구용' if p.get('status')=='exploratory' else '미검증',p.get('sample_n'),p.get('symbol_n'),p.get('quarter_n'),*[p.get(k) for k in ['stop_first_pct','mean_mae_pct','mean_net_return_pct','payoff_ratio','strategy_win_pct','up_pct','down_or_flat_pct','target_first_pct','mean_mfe_pct','distance_mean']],json.dumps(p.get('year_distribution',{}),ensure_ascii=False),p.get('cap500_audit',{}).get('cap500_mean_return_pct'),p.get('cap500_audit',{}).get('all_mean_return_pct')])
    ws=sheet('v22_기술지표',['코드','종목','일봉RSI14','주봉RSI14잠정','상승다이버전스','하락다이버전스','50상향','50하향','과매수','과매도','복합조건','패턴후보'])
    for s in q['stocks']:
        t=s.get('technical',{});ws.append([s['code'],s['name'],*[t.get(k) for k in ['daily_rsi14','weekly_rsi14_partial','bullish_divergence_confirmed','bearish_divergence_confirmed','cross50_up','cross50_down','overbought','oversold','compound']],json.dumps(t.get('patterns',{}),ensure_ascii=False)])
    ws=sheet('v22_공식재무',['코드','업종코드현재','사업연도','연결구분','지표','값','공개시점제한','재무상태'])
    for s in official.get('dart',{}).get('stocks',[]):
        for r in s.get('reports',[]):
            for key,value in r['metrics']['values'].items():ws.append([s['code'],s.get('industry_code_current'),r['year'],r['fs_div'],key,value,'접수일 < 기준일; 역사 백테스트 미반영',s['status']])
    if ws.max_row==1:ws.append(['미확인',None,None,None,None,None,'API 인증키 필요',official['dart']['status']])
    ws=sheet('v22_검증범위',['항목','내용'])
    for key,value in {'기준일':q['asof'],'코드실행':'성공. 데이터 완전성 통과와 다름','원본수집':f"{a['fetched']}/{a['requested']}",'정상연속이력':f"{a['strict_valid']}/{a['requested']}",'독립대조':f"{a['independently_matched']}/{a['requested']}",'v3.0':'진행 보류','가정':'다음 시가 진입, 왕복 0.4%p 비용, 동일 봉 목표/손절은 손절 우선','제외':'거래량0 미래 경로는 체결 미확인으로 제외; 표본선택편향 잔존','편향':'상장폐지/과거구성종목, 기업행위 공식이력, PIT 재무·업종, 독립 예측검증 미완료','RSI비교':'동일 데이터·청산 규칙의 기술통계. 미래 상승확률 또는 RSI 개선효과 검증 아님','공급원':'NAVER fchart / Yahoo chart; 원본과 해시 Actions artifact 보관'}.items():ws.append([key,value])
    for ws in wb:
        if not ws.title.startswith('v22_'):continue
        ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='193B5A');cell.alignment=Alignment(wrap_text=True,vertical='center')
        ws.row_dimensions[1].height=34
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width=24 if len(col)>1 else 18
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value,float):c.number_format='0.00'
        if ws.title=='v22_검증범위':ws.column_dimensions['B'].width=100
    stamp=datetime.now(ZoneInfo('Asia/Seoul')).strftime('%Y%m%d_%H%M%S');filename=f'한국주식_v22_통합검증_{stamp}.xlsx'
    wb.save(reportdir/filename)
    after={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in reportdir.glob('*.xlsx') if p.name in before}
    if before!=after:raise RuntimeError('historical workbook modified')
    q['report_file']='reports/'+filename;q['preserved_report_hashes']=before
    (site/'quant_v22.json').write_text(json.dumps(q,ensure_ascii=False,indent=2,allow_nan=False))
    archive=site/'archive';archive.mkdir(exist_ok=True)
    for p in ['quant_v22.json','reliability.json','official_v22.json']:shutil.copyfile(site/p,archive/(p[:-5]+'_'+stamp+'.json'))
    history=json.loads((site/'history.json').read_text());history.insert(0,{'date':q['asof'],'file':q['report_file']});(site/'history.json').write_text(json.dumps(history,ensure_ascii=False,indent=2))
    print(json.dumps({'report':filename,'preserved_dated_reports':len(before),'v3_ready':False},ensure_ascii=False))
if __name__=='__main__':save()
